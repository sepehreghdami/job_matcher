# services/docx_service.py
"""Extract plain text from a .docx file's raw bytes."""

import io

from docx import Document


def extract_docx_text(data: bytes) -> str:
    """
    Return the concatenated text of all paragraphs and table cells.

    Returns an empty string when nothing can be extracted (corrupt or
    non-.docx file). Callers treat an empty or too-short result as
    "unsupported".
    """
    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    except Exception as e:
        print(f"[docx] extraction failed: {type(e).__name__}: {e}")
        return ""

    return "\n".join(p for p in parts if p.strip()).strip()
